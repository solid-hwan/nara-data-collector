"""Preserve JSON as the canonical export; Excel cells are explicit text."""
import csv
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlsplit
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from .service import canonical


def readable_cell(sheet, value, *, link=False, header=False):
    text = '' if value is None else str(value)
    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]', '\ufffd', text)
    if len(text) > 30000:
        text = text[:29900] + '… [전체 원문: records / long_values]'
    cell = WriteOnlyCell(sheet, value=text)
    cell.data_type = 's'
    if header:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='193A52')
    if link and len(text) <= 2048:
        try:
            parts = urlsplit(text)
            if parts.scheme in ('http', 'https') and parts.hostname:
                cell.hyperlink = text
                cell.style = 'Hyperlink'
        except ValueError:
            pass
    return cell


def export_records(records,destination, *, context=None):
    destination=Path(destination);destination.mkdir(parents=True,exist_ok=True)
    rows=[{'record_index':str(i+1),'raw':canonical(r['raw']),'normalized':canonical(r['normalized']),
           'missing_reasons':canonical(r['missing_reasons'])} for i,r in enumerate(records)]
    columns=['record_index','raw','normalized','missing_reasons']
    with (destination/'records.jsonl').open('w',encoding='utf-8') as out:
        for r in records:out.write(json.dumps(r,ensure_ascii=False)+'\n')
    with (destination/'records.csv').open('w',encoding='utf-8-sig',newline='') as out:
        writer=csv.DictWriter(out,fieldnames=columns);writer.writeheader();writer.writerows(rows)
    workbook=Workbook(write_only=True)
    overview=workbook.create_sheet('Overview');metadata=workbook.create_sheet('Collection')
    sheet=workbook.create_sheet('records');overflow=workbook.create_sheet('long_values')
    headers=['공고번호','차수','공고명','공고기관','수요기관','게시일시','마감일시','추정가격(원문)','원문 링크','결과 의미','원문 행']
    overview.freeze_panes='A2'
    overview.auto_filter.ref=f'A1:K{len(records)+1}'
    for col,width in [('A',22),('B',10),('C',65),('D',30),('E',30),('F',23),('G',23),('H',25),('I',60),('J',30),('K',12)]:
        overview.column_dimensions[col].width=width
    overview.append([readable_cell(overview,x,header=True) for x in headers])
    for index,record in enumerate(records,1):
        raw=record['raw'];normalized=record['normalized']
        values=[raw.get('bidNtceNo') or raw.get('untyCntrctNo'),raw.get('bidNtceOrd') or raw.get('cntrctOrd'),
                normalized.get('title'),raw.get('ntceInsttNm') or raw.get('cntrctInsttNm'),raw.get('dminsttNm'),
                raw.get('bidNtceDt'),raw.get('bidClseDt'),raw.get('presmptPrce'),
                raw.get('bidNtceDtlUrl') or raw.get('bidNtceUrl') or raw.get('cntrctDtlInfoUrl'),
                normalized.get('record_semantics'),index]
        overview.append([readable_cell(overview,x,link=i==8) for i,x in enumerate(values)])
    context=context or {}
    scope={key:context.get(key) for key in ('job_id','status','partial','reason','query','source_total_count',
        'source_total_reason','read_count','filtered_count','unique_count','calls_reserved','max_calls','checkpoint',
        'failed_segments','created_at','updated_at','snapshot_guarantee')}
    scope.update(exported_at=datetime.now(timezone.utc).isoformat(),exported_rows=len(records),
        scope_reason=None if context else 'collection_context_not_provided',
        overview_note='보기용 시트. 빈칸은 원문값 없음. 긴 값은 표시만 축약하며 전체는 records/long_values에 보존.')
    metadata.column_dimensions['A'].width=28;metadata.column_dimensions['B'].width=100
    for key,value in scope.items():
        metadata.append([readable_cell(metadata,key),readable_cell(metadata,canonical(value))])
    def write(ws,values):
        cells=[]
        for value in values:
            cell=WriteOnlyCell(ws,value=str(value));cell.data_type='s';cells.append(cell)
        ws.append(cells)
    write(sheet,columns);write(overflow,['record_index','column','chunk_index','text'])
    for row in rows:
        values=[]
        for key in columns:
            value=row[key]
            if len(value)>32000:
                for n,pos in enumerate(range(0,len(value),32000)):
                    write(overflow,[row['record_index'],key,n,value[pos:pos+32000]])
                value='[long_values:'+row['record_index']+':'+key+']'
            values.append(value)
        write(sheet,values)
    workbook.save(destination/'records.xlsx')
    manifest={'rows':len(rows),'columns':columns,'canonical':'records.jsonl','excel_long_values':'long_values: join text by record_index/column/chunk_index',
              'partial':scope['partial'],'status':scope['status'],'scope_reason':scope['scope_reason'],
              'collection':scope,'overview_sheet':'Overview','collection_sheet':'Collection',
              'amounts_and_identifiers':'preserved as source JSON strings','formula_policy':'all XLSX cells explicit text; CSV contains canonical JSON objects'}
    (destination/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    return manifest
